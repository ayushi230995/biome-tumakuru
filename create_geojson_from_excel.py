"""Build biome_work_tumkur.geojson from the BIOME work Excel file.

Sheet "BIOME Work Details" holds two tables one after the other, each with its
own header row starting with "Date,Sl.No":
  1. Schools / hostels  -> rooftop rainwater harvesting (RWH)
  2. Open wells         -> open well rejuvenation
Sheet "Summary" adds a locality for each school, matched on Sl.No.
Unlike a CSV export, the Excel keeps the hyperlinks behind the picture cells.

Usage: python create_geojson_from_excel.py <workbook.xlsx>
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT_PATH = HERE / 'biome_work_tumkur.geojson'

RWH = 'Rooftop rainwater harvesting'
WELL = 'Open well rejuvenation'


def text(cell):
    return '' if cell.value is None else str(cell.value).strip()


def link(cell):
    """The cell's hyperlink, or its text if that is itself a URL."""
    url = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else text(cell)
    return url if url.startswith('http') else None


def number(cell):
    """'400 SQM' -> 400, '11000 Ltr' -> 11000, 'No' / 'Nil' / '' / 0 -> None."""
    nums = re.findall(r'\d+(?:\.\d+)?', text(cell).replace(',', ''))
    if len(nums) != 1 or float(nums[0]) == 0:
        return None
    n = float(nums[0])
    return int(n) if n.is_integer() else n


def litres_range(cell):
    """'2800 Ltr' -> 2800, '20000 to 25000 Ltr' -> [20000, 25000]."""
    nums = [int(n) for n in re.findall(r'\d+', text(cell).replace(',', ''))]
    return nums[0] if len(nums) == 1 else nums or None


def year(cell):
    """'2023 - 2024' / '2023-2024' -> '2023-24'."""
    start, end = re.findall(r'\d{4}', text(cell))
    return f'{start}-{end[2:]}'


def interventions(cell):
    """'Roof top RWH, Recharge well, Storage stump' -> ['Roof top RWH', 'Recharge well', 'Storage sump']."""
    value = re.sub(r'^Open well rejuvenation\s*-\s*', '', text(cell))
    value = re.sub(r'\bstump\b', 'sump', value)  # typo in source sheet
    return [part.strip()[:1].upper() + part.strip()[1:] for part in value.split(',') if part.strip()]


def split_tables(ws):
    """Yield one list of {header: cell} dicts per table on the sheet.

    A trailing unit in a header is dropped ('Well Depth (Feet)' -> 'Well Depth'),
    so adding units to column names in the sheet doesn't break the lookups below.
    """
    header, rows = None, []
    for cells in ws.iter_rows():
        values = [text(c) for c in cells]
        if values[0] == 'Date':
            if rows:
                yield rows
            header, rows = [re.sub(r'\s*\([^)]*\)$', '', v) for v in values], []
        elif header and any(values):
            rows.append(dict(zip(header, cells)))
    if rows:
        yield rows


def localities(ws):
    """Sl.No -> 'Location' from the Summary sheet."""
    found, header = {}, None
    for cells in ws.iter_rows():
        values = [text(c) for c in cells]
        if 'Sl.No' in values and 'Location' in values:
            header = values
        elif header:
            row = dict(zip(header, values))
            if row['Sl.No'].isdigit() and row['Location']:
                found[int(row['Sl.No'])] = row['Location']
    return found


def rwh_properties(row, locality):
    name = text(row['School Name'])
    return {
        'category': RWH,
        'institution': 'Hostel' if 'hostel' in name.lower() else 'School',
        'name': name,
        'locality': locality,
        'year': year(row['Date']),
        'funder': text(row['Funded By']),
        'interventions': interventions(row['What we did']),
        'roof_area_m2': number(row['Building Roof Area covered']),
        'sump_l': number(row['RWH Sump Size']),
        'rain_barrel_l': number(row['Rain Barrel Ltr']),
        'students': number(row['Students']),
        'teaching_staff': number(row['Teaching Staff']),
        'daily_water_demand_l': litres_range(row['Daily Water Demand']),
        'potential_harvest_l': number(row['Potential Rainwater harvested']),
        'rainfall_mm': number(row['Rainfall MM']) if 'Rainfall MM' in row else None,
        'maps_url': link(row['Locations']),
        'photos_url': link(row['Link - Pictures']),
    }


def well_properties(row, locality):
    return {
        'category': WELL,
        'institution': 'Open well',
        'name': text(row['Well Name']),
        'locality': locality,
        'year': year(row['Date']),
        'funder': text(row['Funded By']),
        'interventions': interventions(row['What we did']),
        'well_diameter_ft': number(row['Well dimentions']),
        'well_depth_ft': number(row['Well Depth']),
        'well_volume_l': number(row['Well Volume']),
        'households': number(row['Households using the well']),
        'maps_url': link(row['Locations']),
        'photos_url': link(row['Link']),
    }


if len(sys.argv) != 2:
    sys.exit(__doc__)
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
rwh_rows, well_rows = split_tables(wb['BIOME Work Details'])
school_localities = localities(wb['Summary'])

features = []
for prefix, rows, to_properties, places in [
    ('rwh', rwh_rows, rwh_properties, school_localities),
    ('well', well_rows, well_properties, {}),
]:
    for row in rows:
        sl = int(text(row['Sl.No']))
        lat, lon = float(text(row['Latitude'])), float(text(row['Longitude']))
        features.append({
            'type': 'Feature',
            'geometry': {'type': 'Point', 'coordinates': [lon, lat]},
            'properties': {'id': f'{prefix}-{sl:02d}', **to_properties(row, places.get(sl))},
        })

with open(OUT_PATH, 'w', encoding='utf-8') as f:
    json.dump({'type': 'FeatureCollection', 'features': features}, f, ensure_ascii=False, indent=1)

with_photos = sum(1 for f in features if f['properties']['photos_url'])
print(f'Wrote {len(features)} sites ({len(rwh_rows)} RWH, {len(well_rows)} wells, '
      f'{with_photos} with photo links) -> {OUT_PATH.name}')
